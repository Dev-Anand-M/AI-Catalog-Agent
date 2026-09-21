import sys
import subprocess

try:
    import docx
except ImportError:
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx"])
    import docx

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="CCCCCC", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def make_styled_docx():
    doc = Document()
    
    # Page setup - Margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
        
    # Styles
    # Title / Main Header
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = p_title.add_run("FINAL YEAR PROJECT & STARTUP REVIEW WORKBOOK\n")
    run_sub.font.name = 'Calibri'
    run_sub.font.size = Pt(12)
    run_sub.font.bold = True
    run_sub.font.color.rgb = RGBColor(100, 110, 120)
    
    run_main = p_title.add_run("ONE-PERSON EXTREME PITCH\n")
    run_main.font.name = 'Calibri'
    run_main.font.size = Pt(26)
    run_main.font.bold = True
    run_main.font.color.rgb = RGBColor(20, 50, 90) # Dark Blue
    
    run_tag = p_title.add_run("Digital Catalog Agent — Filled College & Startup Review Edition\n")
    run_tag.font.name = 'Calibri'
    run_tag.font.size = Pt(14)
    run_tag.font.bold = True
    run_tag.font.color.rgb = RGBColor(0, 128, 128) # Teal
    
    run_meta = p_title.add_run("Includes Step-by-Step Solo Protocol, Filled 6-Stakeholder Challenge Log & Final Verdict\n")
    run_meta.font.name = 'Calibri'
    run_meta.font.size = Pt(10)
    run_meta.font.italic = True
    run_meta.font.color.rgb = RGBColor(120, 120, 120)

    doc.add_paragraph() # Spacing

    # Philosophy Box
    p_phil = doc.add_paragraph()
    p_phil.paragraph_format.space_before = Pt(6)
    p_phil.paragraph_format.space_after = Pt(12)
    
    # Create callout box using 1x1 table
    tbl_callout = doc.add_table(rows=1, cols=1)
    tbl_callout.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_callout.autofit = False
    cell_c = tbl_callout.rows[0].cells[0]
    cell_c.width = Inches(6.8)
    set_cell_background(cell_c, "F0F4F8") # Very light blue-gray
    set_cell_margins(cell_c, top=140, bottom=140, left=200, right=200)
    
    cp = cell_c.paragraphs[0]
    r1 = cp.add_run("[PHILOSOPHY] WHAT IS THE ONE-PERSON EXTREME PITCH?\n")
    r1.font.name = 'Calibri'
    r1.font.bold = True
    r1.font.size = Pt(11)
    r1.font.color.rgb = RGBColor(20, 50, 90)
    
    r2 = cp.add_run('"Be your own toughest examiner before real-world users test you."\n')
    r2.font.name = 'Calibri'
    r2.font.bold = True
    r2.font.italic = True
    r2.font.size = Pt(10.5)
    r2.font.color.rgb = RGBColor(180, 40, 40) # Crimson
    
    r3 = cp.add_run("Extreme Pitch is NOT created to judge or fail you. Its goal is to help you find weak assumptions, security bugs, scaling bottlenecks, and UX issues on paper BEFORE your external college examiners or real users find them! When teammates are unavailable, the Solo Multi-Hat Protocol maintains 100% of the 3-man pitch efficiency.")
    r3.font.name = 'Calibri'
    r3.font.size = Pt(9.5)
    r3.font.color.rgb = RGBColor(60, 60, 60)

    doc.add_paragraph()

    # Part 1: How It Works
    h1 = doc.add_heading(level=1)
    r_h1 = h1.add_run("PART 1: HOW THE ONE-PERSON EXTREME PITCH WORKS")
    r_h1.font.name = 'Calibri'
    r_h1.font.color.rgb = RGBColor(20, 50, 90)
    r_h1.font.bold = True

    p_guide = doc.add_paragraph()
    p_guide.add_run("Instead of 3 physical participants, you operate in a timed ").font.name = 'Calibri'
    r_b = p_guide.add_run("Multi-Hat Split-Persona Protocol")
    r_b.bold = True
    r_b.font.name = 'Calibri'
    p_guide.add_run(". You cycle through 3 distinct operational mindsets during a 60-minute review session:\n").font.name = 'Calibri'

    bullet1 = doc.add_paragraph(style='List Bullet')
    r = bullet1.add_run("Hat 1: Product Owner (Defender) — ")
    r.bold = True
    r.font.name = 'Calibri'
    bullet1.add_run("Presents core architecture, data flows, and tech stack; defends logic without defensiveness.").font.name = 'Calibri'

    bullet2 = doc.add_paragraph(style='List Bullet')
    r = bullet2.add_run("Hat 2: Adversarial Challenger (Red Team) — ")
    r.bold = True
    r.font.name = 'Calibri'
    bullet2.add_run("Draws 4 to 6 Stakeholder Role Cards (Security, Cloud Arch, Investor, QA, Privacy, Legal) and attacks the system.").font.name = 'Calibri'

    bullet3 = doc.add_paragraph(style='List Bullet')
    r = bullet3.add_run("Hat 3: Objective Auditor (Facilitator) — ")
    r.bold = True
    r.font.name = 'Calibri'
    bullet3.add_run("Logs signals, fills 9-point scorecard objectively, maps single-owner action items, and renders final verdict.").font.name = 'Calibri'

    doc.add_paragraph()

    # Helper function to build styled tables
    def add_custom_table(headers, rows_data, col_widths=None, header_bg="1F4E79", alt_bg="F9FBFD"):
        table = doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        set_table_borders(table, color="D3D3D3", sz="4")
        
        # Header Row
        hdr_row = table.rows[0]
        for idx, text in enumerate(headers):
            cell = hdr_row.cells[idx]
            if col_widths and idx < len(col_widths):
                cell.width = Inches(col_widths[idx])
            set_cell_background(cell, header_bg)
            set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(text)
            r.font.name = 'Calibri'
            r.font.bold = True
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(255, 255, 255)
            
        # Data Rows
        for r_idx, row_values in enumerate(rows_data):
            row = table.rows[r_idx + 1]
            bg = alt_bg if r_idx % 2 == 1 else "FFFFFF"
            for c_idx, val in enumerate(row_values):
                cell = row.cells[c_idx]
                if col_widths and c_idx < len(col_widths):
                    cell.width = Inches(col_widths[c_idx])
                set_cell_background(cell, bg)
                set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                r = p.add_run(str(val))
                r.font.name = 'Calibri'
                r.font.size = Pt(9.0)
                r.font.color.rgb = RGBColor(40, 40, 40)
                
                # Highlight status badges
                if str(val) in ["ACCEPTED", "APPROVED"]:
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(30, 130, 60) # Green
                elif str(val) in ["NEEDS IMPROVEMENT", "APPROVED WITH CHANGES"]:
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(180, 100, 0) # Orange
                elif str(val) in ["RESEARCH REQUIRED"]:
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(0, 100, 180) # Blue
                elif str(val) in ["P1", "HIGH"]:
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(180, 40, 40) # Red
                    
        doc.add_paragraph() # Spacing

    # Section Headers & Tables
    # Section 1
    h_s1 = doc.add_heading(level=2)
    r = h_s1.add_run("SECTION 1: SOLO SESSION INFORMATION")
    r.font.color.rgb = RGBColor(20, 50, 90)
    
    s1_headers = ["Field", "Value / Session Details"]
    s1_data = [
        ["Project Name", "Digital Catalog Agent (AI Catalog Agent)"],
        ["Date & Time", "2026-07-26 | 17:30 IST"],
        ["Version / Tag", "v1.0 (Final Year Project & Startup Review Build)"],
        ["Duration", "60 Minutes"],
        ["Solo Presenter & Reviewer", "Dev Anand (Lead Full-Stack Developer & Solo Auditor)"],
        ["Project Team Members", "Dev Anand, Hemanth, Dhanushree"],
        ["Active Stakeholder Cards", "1. Security Engineer  2. Cloud Architect  3. Investor\n4. QA & Resilience Tester  5. Privacy Officer  6. Compliance Lawyer"]
    ]
    add_custom_table(s1_headers, s1_data, [2.2, 4.6], header_bg="1F4E79")

    # Section 2
    h_s2 = doc.add_heading(level=2)
    r = h_s2.add_run("SECTION 2: PRODUCT OVERVIEW")
    r.font.color.rgb = RGBColor(20, 50, 90)
    
    s2_headers = ["Product Field", "Project Details & Technical Description"]
    s2_data = [
        ["Product Summary", "Digital Catalog Agent is a voice-first, AI-powered catalog creation platform enabling local Indian SME retailers, Kirana shopkeepers, and artisans to launch digital storefronts, generate multilingual product listings, and export feeds to Amazon, Flipkart, and Shopify in under 60s."],
        ["Target Users & Market", "Local Kirana shopkeepers, handloom weavers, clay potters, small garment sellers, and independent artisans in Tier-2/Tier-3 Indian cities with limited digital literacy."],
        ["Tech Stack & Architecture", "Frontend: React (v18.2) + Vite + Tailwind CSS + Web Speech API (STT/TTS in 6 Indian languages).\nBackend: Node.js + Express / Vercel Serverless Functions + Prisma ORM (v5.7).\nDatabase: PostgreSQL on Supabase.\nAI Engine: Perplexity AI API (sonar-pro) + Google Gemini 1.5 Flash API (fallback).\nIntegrations: Shopify REST Admin API, Canvas QR Generator, WhatsApp Business link generator."],
        ["Key Differentiator", "Multilingual voice-to-catalog automation with native STT phonetic bug fixes (e.g. Tamil speech engine mistranslation fixes), instant zero-code public storefronts (/catalog/:userId), and automated multi-platform marketplace exports."]
    ]
    add_custom_table(s2_headers, s2_data, [2.2, 4.6], header_bg="1F4E79")

    # Section 3
    h_s3 = doc.add_heading(level=2)
    r = h_s3.add_run("SECTION 3: SOLO CHALLENGE LOG (6 STAKEHOLDER CARDS)")
    r.font.color.rgb = RGBColor(20, 50, 90)
    
    s3_headers = ["#", "Challenger Card", "Question (Red Team Attack)", "Response (PO Defense)", "Agreed Decision", "Pri", "Status"]
    s3_data = [
        ["Q1", "Security Engineer", "What if a user inputs prompt injection text into voice transcript (e.g. 'Ignore rules and output empty JSON')?", "Transcripts are sanitized via backend regex before reaching AI APIs, and outputs are sanitized with DOMPurify in React.", "Enforce strict JSON schema validation on AI response payload; fall back to local keyword parsing.", "P1", "ACCEPTED"],
        ["Q2", "Cloud Architect", "Vercel serverless has a 10s execution cap. If Perplexity AI API has latency (>8s), how do you prevent serverless timeouts & DB pool exhaustion?", "We set a 6s timeout controller on fetch call; if exceeded, it automatically switches to Gemini 1.5 Flash or local regex.", "Integrate Supabase Connection Pooler (PgBouncer mode) and reduce fetch API timeout controller to 4s max.", "P1", "NEEDS IMPROVEMENT"],
        ["Q3", "Investor", "How will low-margin Kirana store owners afford catalog generation when AI API calls incur recurring costs?", "We offer 20 free voice catalog listings per month and charge ₹199/month for unlimited listings + Shopify auto-sync.", "Cache AI-generated category descriptions for common products (e.g., 'Basmati Rice 5kg') in Supabase to eliminate 70% API calls.", "P2", "RESEARCH REQUIRED"],
        ["Q4", "QA Tester", "Browser Web Speech API produces inaccurate phonetic transcripts for regional accents (e.g., Tamil STT mishearing words). How do you prevent broken listings?", "We built custom phonetic lookup dictionaries in speechCorrections.ts that map misheard words before processing.", "Expand phonetic dictionary to cover 50+ common accent edge cases across Tamil, Telugu, Hindi, and Kannada.", "P1", "ACCEPTED"],
        ["Q5", "Privacy Officer", "Are raw voice recordings or user phone numbers transmitted to external AI servers during voice catalog creation?", "Web Speech API processes speech locally in browser; only plain text transcript is sent. No audio or PII saved.", "Add explicit privacy badge on voice input modal confirming zero audio retention or external PII storage.", "P2", "ACCEPTED"],
        ["Q6", "Compliance Lawyer", "If export to Amazon/Flipkart contains mismatched GST or missing legal declarations, who is liable for listing rejection?", "App auto-populates standard Indian GST tiers (5%, 12%, 18%) and sets default Country of Origin to 'India'.", "Require seller to review and confirm a 'Mandatory Legal & Tax Declaration' modal before executing feed exports.", "P1", "ACCEPTED"]
    ]
    add_custom_table(s3_headers, s3_data, [0.4, 1.1, 1.6, 1.5, 1.4, 0.4, 0.9], header_bg="1F4E79")

    # Section 4 & 5
    h_s4 = doc.add_heading(level=2)
    r = h_s4.add_run("SECTION 4: KEY INSIGHTS")
    r.font.color.rgb = RGBColor(20, 50, 90)
    
    s4_headers = ["#", "Category", "Core Discovery / Insight", "Strategic Implication", "Severity"]
    s4_data = [
        ["INS-01", "Security & AI", "Relying solely on LLM prompt instructions is unsafe; backend JSON schema validation and deterministic fallbacks are essential.", "All AI API responses must be validated against a JSON Schema guardrail before being written to PostgreSQL.", "HIGH"],
        ["INS-02", "Infrastructure & DB", "Serverless deployments require aggressive API timeouts (4s max) and connection pooling to prevent cascading failures.", "Configure PgBouncer connection pooling on Supabase to prevent connection exhaustion during traffic bursts.", "HIGH"],
        ["INS-03", "UX & Compliance", "E-commerce catalog generators must enforce a human-in-the-loop verification step before pushing listings to external channels.", "Enforce mandatory preview and GST confirmation modals prior to triggering bulk marketplace export downloads.", "MEDIUM"]
    ]
    add_custom_table(s4_headers, s4_data, [0.6, 1.3, 2.2, 2.2, 0.8], header_bg="1F4E79")

    h_s5 = doc.add_heading(level=2)
    r = h_s5.add_run("SECTION 5: RECOMMENDATIONS")
    r.font.color.rgb = RGBColor(20, 50, 90)
    
    s5_headers = ["ID", "Pri", "Recommendation Proposal", "Justification / Reason", "Origin", "Impact", "Diff", "Status"]
    s5_data = [
        ["REC-01", "P1", "Implement backend JSON schema validator & sanitize AI output payload.", "Prevents prompt injection & corrupted JSON structures from crashing frontend.", "Security Eng", "High", "Med", "ACCEPTED"],
        ["REC-02", "P1", "Configure Supabase PgBouncer connection pooler & reduce API timeout to 4s.", "Prevents serverless execution timeouts and database pool exhaustion.", "Cloud Arch", "High", "Med", "ACCEPTED"],
        ["REC-03", "P1", "Add mandatory Seller Review & GST Sign-Off Modal prior to exports.", "Protects platform and seller from marketplace listing rejections.", "Lawyer", "High", "Low", "ACCEPTED"],
        ["REC-04", "P1", "Expand speechCorrections.ts phonetic dictionary to cover 50+ speech errors.", "Guarantees high listing accuracy across diverse Indian accents (Hindi, Tamil, Telugu).", "QA Tester", "High", "Med", "ACCEPTED"]
    ]
    add_custom_table(s5_headers, s5_data, [0.6, 0.4, 1.8, 1.8, 0.9, 0.5, 0.4, 0.9], header_bg="1F4E79")

    # Section 6
    h_s6 = doc.add_heading(level=2)
    r = h_s6.add_run("SECTION 6: ACTION ITEMS")
    r.font.color.rgb = RGBColor(20, 50, 90)
    
    s6_headers = ["Task ID", "Action Task Description", "Single Owner", "Deadline", "Priority", "Dependencies", "Status"]
    s6_data = [
        ["ACT-01", "Build JSON schema validator in server/controllers/aiController.js", "Dev Anand", "2026-07-28", "P1", "None", "ACCEPTED"],
        ["ACT-02", "Enable PgBouncer connection string in prisma/schema.prisma & Supabase", "Dev Anand", "2026-07-29", "P1", "ACT-01", "ACCEPTED"],
        ["ACT-03", "Create Seller Review Modal in client/src/components/ExportModal.jsx", "Dev Anand", "2026-07-30", "P1", "None", "ACCEPTED"],
        ["ACT-04", "Add 30 additional regional phonetic mappings in speechCorrections.js", "Dev Anand", "2026-07-31", "P1", "None", "ACCEPTED"]
    ]
    add_custom_table(s6_headers, s6_data, [0.7, 2.3, 0.9, 0.9, 0.5, 0.8, 0.9], header_bg="1F4E79")

    # Section 7
    h_s7 = doc.add_heading(level=2)
    r = h_s7.add_run("SECTION 7: SOLO AUDIT SCORECARD")
    r.font.color.rgb = RGBColor(20, 50, 90)
    
    s7_headers = ["Evaluation Criteria (1.0 - 10.0 Scale)", "Solo Score", "Rationale & Friction Check"]
    s7_data = [
        ["1. Problem Understanding (Domain & pain points)", "9.50", "Deep understanding of digital exclusion barriers for Indian SME retailers."],
        ["2. Solution Clarity (Architecture & elegance)", "9.00", "Clean hybrid architecture (React + Express + Supabase + Perplexity/Gemini)."],
        ["3. Technical Knowledge (Security, DB, scaling, infra)", "8.75", "Robust voice processing; identified serverless DB pool limits under load."],
        ["4. Business Thinking (Monetization & unit economics)", "8.00", "Good freemium model; needs caching strategy to optimize per-request API costs."],
        ["5. Communication (Structure & clarity)", "9.25", "Clear documentation, thorough project report, and explicit API contracts."],
        ["6. Confidence & Composure (Poise under attack)", "9.00", "Accepted technical blindspots (STT errors, serverless timeouts) without excuses."],
        ["7. Handling Blindspots (Edge case navigation)", "8.75", "Resolved voice phonetic misspellings and multi-language fallback logic."],
        ["8. Openness To Feedback (Willingness to adapt arch)", "9.25", "Readily adopted JSON schema validation, PgBouncer pooling, and export sign-off."],
        ["9. Overall Pitch Execution (Rigor & professionalism)", "9.00", "High-impact pitch defense with complete remediation roadmap and clean code."],
        ["GRAND SCORECARD AVERAGE", "8.94 / 10", "OUTSTANDING TECHNICAL DEFENSE"]
    ]
    add_custom_table(s7_headers, s7_data, [3.2, 0.9, 2.7], header_bg="1F4E79")

    # Section 8
    h_s8 = doc.add_heading(level=2)
    r = h_s8.add_run("SECTION 8: SESSION OUTCOME & FINAL VERDICT")
    r.font.color.rgb = RGBColor(20, 50, 90)
    
    s8_headers = ["Review Metric", "Summary & Final Verdict Details"]
    s8_data = [
        ["Biggest Technical Strength", "Multilingual voice-to-catalog pipeline with real-time phonetic correction (speechCorrections.ts) and zero-downtime dual AI fallback architecture (Perplexity + Gemini)."],
        ["Biggest Vulnerability", "Potential serverless function timeouts on Vercel if primary AI API latency spikes during peak user traffic."],
        ["Most Valuable Insight", "AI-assisted voice cataloging must pair natural language input with deterministic regex fallbacks, strict JSON schema validation, and human sign-off before marketplace exports."],
        ["Key Lessons Learned", "1. Regional speech engine errors must be corrected locally before sending transcripts to AI APIs.\n2. Serverless database connections must use connection pooling (PgBouncer) to withstand concurrency spikes."],
        ["Final Recommendation", "Project is fully approved for final submission upon executing Action Items ACT-01 (JSON Schema Validation) and ACT-02 (PgBouncer Pooling)."],
        ["FINAL SESSION VERDICT", "APPROVED WITH CHANGES"]
    ]
    add_custom_table(s8_headers, s8_data, [2.2, 4.6], header_bg="1F4E79")

    # Save output docx
    output_path = "FILLED_EXTREME_PITCH_AI_CATALOG_AGENT.docx"
    doc.save(output_path)
    print(f"Successfully generated styled DOCX: {output_path}")

if __name__ == "__main__":
    make_styled_docx()
