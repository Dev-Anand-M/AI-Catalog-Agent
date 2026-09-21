import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

doc = Document()

# Page Margins
for section in doc.sections:
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

# Banner Header
header_table = doc.add_table(rows=1, cols=1)
header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
cell = header_table.cell(0, 0)
set_cell_background(cell, "0F172A")
set_cell_margins(cell, top=200, bottom=200, left=200, right=200)

p = cell.paragraphs[0]
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("⚡ WEEK 1 CONTRIBUTION LOG — NO PUSH")
run.font.name = 'Calibri'
run.font.size = Pt(22)
run.font.bold = True
run.font.color.rgb = RGBColor(56, 189, 248)

p2 = cell.add_paragraph()
p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
run2 = p2.add_run("DEVELOPER: DEV ANAND  |  ROLE: TECH LEAD  |  WEEK 1 (AUG 5 - AUG 12)")
run2.font.name = 'Calibri'
run2.font.size = Pt(11)
run2.font.bold = True
run2.font.color.rgb = RGBColor(253, 224, 71)

doc.add_paragraph()

# Section 1: Executive Overview
h1 = doc.add_heading("📌 1. EXECUTIVE SUMMARY & ASSIGNED SCOPE", level=2)
h1.runs[0].font.name = 'Calibri'
h1.runs[0].font.color.rgb = RGBColor(30, 64, 175)

p = doc.add_paragraph()
run = p.add_run("As Tech Lead for Week 1 (Core Storefront Stage), Dev Anand led the foundational backend architecture, Express API router setup, Supabase database integration, and voice-to-catalog STT generation pipeline.")
run.font.name = 'Calibri'
run.font.size = Pt(11)

# Section 2: Codebase File Changes
h2 = doc.add_heading("💻 2. ARCHITECTURAL FILE IMPLEMENTATIONS & CODE PROOF", level=2)
h2.runs[0].font.name = 'Calibri'
h2.runs[0].font.color.rgb = RGBColor(30, 64, 175)

table = doc.add_table(rows=1, cols=3)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
hdr_cells = table.rows[0].cells
headers = ["File Path", "Feature Module", "Implementation Details"]
widths = [Inches(2.2), Inches(1.8), Inches(2.8)]

for i, title in enumerate(headers):
    hdr_cells[i].text = title
    set_cell_background(hdr_cells[i], "1E3A8A")
    set_cell_margins(hdr_cells[i], top=120, bottom=120, left=150, right=150)
    p = hdr_cells[i].paragraphs[0]
    p.runs[0].font.name = 'Calibri'
    p.runs[0].font.bold = True
    p.runs[0].font.size = Pt(10)
    p.runs[0].font.color.rgb = RGBColor(255, 255, 255)

impls = [
    ("server/src/routes/ai.js", "Voice Input & AI Generation", "Added action query re-writer middleware for ?action=generate & ?action=enhance endpoints. Integrated Gemini 1.5 Flash STT payload parsing for 6 Indian languages."),
    ("server/src/routes/auth.js", "Auth API Router", "Implemented ?action=signup, ?action=login, and ?action=me routing rules with bcrypt password hashing (10 salt rounds) and JWT token generation."),
    ("server/src/db.js", "Supabase REST Driver", "Built lightweight Supabase REST database driver replacing Prisma ORM to connect via SUPABASE_URL and SUPABASE_SERVICE_ROLE_KEY."),
    ("server/src/index.js", "Express Server Entry", "Configured Express app middleware, CORS headers, error handlers, and HTTP port 3001 listener.")
]

for idx, (fpath, module_name, desc) in enumerate(impls):
    row_cells = table.add_row().cells
    bg = "F8FAFC" if idx % 2 == 1 else "FFFFFF"
    for i, val in enumerate([fpath, module_name, desc]):
        row_cells[i].text = val
        set_cell_background(row_cells[i], bg)
        set_cell_margins(row_cells[i], top=100, bottom=100, left=120, right=120)
        p = row_cells[i].paragraphs[0]
        p.runs[0].font.name = 'Calibri'
        p.runs[0].font.size = Pt(10)
        if i == 0:
            p.runs[0].font.bold = True

doc.add_paragraph()

# Section 3: Empirical Verification & Bug Fixes
h3 = doc.add_heading("🔧 3. EMPIRICAL VERIFICATION & BUG RESOLUTION LOG", level=2)
h3.runs[0].font.name = 'Calibri'
h3.runs[0].font.color.rgb = RGBColor(30, 64, 175)

verifs = [
    ("Fixed 404 Route Mismatch", "Resolved query-parameter route lookup by introducing URL re-writer middleware mapping ?action=signup directly to router endpoints."),
    ("Resolved HTTP 500 Prisma Error", "Fixed database initialization failure by building server/src/db.js REST driver connecting straight to live Supabase Cloud PostgreSQL."),
    ("Verified Auth Signup Endpoint", "Tested POST http://localhost:3001/api/auth?action=signup — Verified HTTP 201 Created and JWT token issuance."),
    ("Verified AI Voice Generator", "Tested POST http://localhost:3001/api/ai?action=generate with audio transcript — Verified HTTP 200 JSON product detail payload.")
]

for title, desc in verifs:
    p = doc.add_paragraph(style='List Bullet')
    r1 = p.add_run(f"{title}: ")
    r1.font.name = 'Calibri'
    r1.font.bold = True
    r1.font.size = Pt(11)
    r2 = p.add_run(desc)
    r2.font.name = 'Calibri'
    r2.font.size = Pt(11)

doc.add_paragraph()

# Section 4: Self-Evaluation Scorecard
h4 = doc.add_heading("⭐ 4. WEEK 1 RUBRIC SELF-EVALUATION SCORECARD", level=2)
h4.runs[0].font.name = 'Calibri'
h4.runs[0].font.color.rgb = RGBColor(30, 64, 175)

score_table = doc.add_table(rows=1, cols=1)
score_cell = score_table.cell(0, 0)
set_cell_background(score_cell, "F0FDF4")
set_cell_margins(score_cell, top=150, bottom=150, left=150, right=150)

scores = [
    "Task Ownership (10 / 10): Completed full API router and database driver pipeline on schedule.",
    "Code Quality & Mobile UX (10 / 10): Clean modular Express code, error boundaries, no secret leaks.",
    "Demo & Understanding (10 / 10): Fully prepared to demonstrate API endpoints during Sunday demo.",
    "Autonomy & Integration (10 / 10): Solved Prisma 500 and 404 router bugs independently without blocking teammates."
]

for idx, item in enumerate(scores):
    p = score_cell.paragraphs[0] if idx == 0 else score_cell.add_paragraph()
    run = p.add_run(f"✓  {item}")
    run.font.name = 'Calibri'
    run.font.size = Pt(10.5)
    run.font.bold = True
    run.font.color.rgb = RGBColor(22, 101, 52)

output_file = r"c:\Users\p3ace\OneDrive\Documents\New folder\AI-Catalog-Agent\DEV_ANAND_WEEK1_CONTRIBUTIONS.docx"
doc.save(output_file)
print(f"Generated: {output_file}")
