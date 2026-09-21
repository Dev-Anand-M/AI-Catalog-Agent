import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="CCCCCC", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def format_run(run, font_name="Calibri", font_size=10.5, bold=False, italic=False, color_rgb=(51,51,51)):
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(*color_rgb)

# -------------------------------------------------------------
# DOCUMENT 1: TEAM ROADMAP AND LEARNING PLAN
# -------------------------------------------------------------
def build_roadmap_docx(filepath):
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("🎴 WELCOME TO THE TAROT CLUB LEVEL-UP CHALLENGE! 🚀")
    format_run(r_title, "Segoe UI", 18, bold=True, color_rgb=(103, 58, 183))

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("AI Catalog Agent — 8-Week Flexible Feature Roadmap & Innovation Plan\nTeam: The Tarot Club  |  Lead: Dev Anand")
    format_run(r_sub, "Calibri", 11, italic=True, color_rgb=(100, 100, 100))

    doc.add_paragraph()

    # Callout Box: Game Plan
    tbl_gp = doc.add_table(rows=1, cols=1)
    tbl_gp.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell_gp = tbl_gp.cell(0, 0)
    set_cell_background(cell_gp, "F3E5F5")
    set_cell_margins(cell_gp, top=120, bottom=120, left=180, right=180)
    
    p_gp_head = cell_gp.paragraphs[0]
    r_gp_head = p_gp_head.add_run("⚡ THE GAME PLAN (College-Friendly Flexible Schedule)")
    format_run(r_gp_head, "Segoe UI", 12, bold=True, color_rgb=(74, 20, 140))
    
    p_gp_body = cell_gp.add_paragraph()
    r_gp_body = p_gp_body.add_run(
        "• MON - THU: 100% Chill & Async. Do micro-tasks or design reading whenever free around college.\n"
        "• SAT - SUN: Weekend Coding Sprint! Build features, pair-review code, and integrate.\n"
        "• SUN EVENING: 15-Minute Demo Showcase & Weekly Evaluation Scorecard Logging!"
    )
    format_run(r_gp_body, "Calibri", 10.5, color_rgb=(51, 51, 51))

    doc.add_paragraph()

    # Section 1: Master Cheat Sheet Table
    p_h1 = doc.add_paragraph()
    r_h1 = p_h1.add_run("📊 MASTER CHEAT SHEET (AT-A-GLANCE ROADMAP)")
    format_run(r_h1, "Segoe UI", 14, bold=True, color_rgb=(103, 58, 183))

    table_data = [
        ["Week", "Quest Stage", "💸 Hemanth (Finance Lead)", "🎨 Dhanushree (UI/UX Code Lead)", "⚡ Dev Anand (Tech Lead)"],
        ["W1", "Core Storefront", "UPI Settings & Bank Details\n(PaymentSettings.jsx)", "Public Catalog Layout & Grid\n(PublicCatalog.jsx)", "Voice input & API routes setup"],
        ["W2", "Payments & i18n", "Dynamic QR Generator & WhatsApp links", "6-Language Regional Keyboard & Selector", "AI response parsing & Supabase DB"],
        ["W3", "Exporter & Prices", "Pricing data structures & API update handlers", "Export Catalog UI cards (Shopify/Amazon)", "Shopify REST sync endpoints"],
        ["W4", "🃏 TAROTHON #1", "💡 NEW FEATURE INNOVATION: Pitch & build!", "💡 NEW FEATURE INNOVATION: Pitch & build!", "Technical mentor & PR reviews"],
        ["W5", "Dashboards", "Payment status toggles & summary UI", "Dashboard layout (Dashboard.jsx) & filters", "Voice command interpreter endpoint"],
        ["W6", "Multimodal UI", "Product edit pricing fields (EditProduct.jsx)", "Camera Capture modal UI (CameraCapture.jsx)", "AI fallback handlers (Gemini/Perplexity)"],
        ["W7", "Landing & Polish", "Payment security review & setup guide polish", "Landing Page UI (Landing.jsx) & docs polish", "Vercel production build & deploy"],
        ["W8", "🏆 TAROTHON #2", "🚀 FINAL SHOWCASE INNOVATION: Build & launch!", "🚀 FINAL SHOWCASE INNOVATION: Build & launch!", "Final pitch defense & launch prep"]
    ]

    tbl_cheat = doc.add_table(rows=len(table_data), cols=5)
    tbl_cheat.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_cheat, color="D1C4E9", sz="6")
    col_widths = [Inches(0.6), Inches(1.3), Inches(1.8), Inches(1.8), Inches(1.5)]

    for row_idx, row_content in enumerate(table_data):
        for col_idx, text in enumerate(row_content):
            cell = tbl_cheat.cell(row_idx, col_idx)
            cell.width = col_widths[col_idx]
            set_cell_margins(cell, top=80, bottom=80, left=80, right=80)
            p = cell.paragraphs[0]
            r = p.add_run(text)
            if row_idx == 0:
                set_cell_background(cell, "512DA8")
                format_run(r, "Segoe UI", 9.5, bold=True, color_rgb=(255, 255, 255))
            elif "TAROTHON" in text:
                set_cell_background(cell, "EDE7F6")
                format_run(r, "Calibri", 9.5, bold=True, color_rgb=(103, 58, 183))
            else:
                bg = "FAFAFA" if row_idx % 2 == 0 else "FFFFFF"
                set_cell_background(cell, bg)
                format_run(r, "Calibri", 9, color_rgb=(51, 51, 51))

    doc.add_paragraph()

    # Section 2: Rubric & PR Checklist
    p_h2 = doc.add_paragraph()
    r_h2 = p_h2.add_run("📝 WEEKLY EVALUATION RUBRIC & PR CHECKLIST")
    format_run(r_h2, "Segoe UI", 14, bold=True, color_rgb=(103, 58, 183))

    rubric_items = [
        "1. Feature Completeness (10 pts): Clean feature implementation and full code integration.",
        "2. Code Quality & Mobile Responsiveness (10 pts): Standard React/Tailwind/Node compliance, no console errors.",
        "3. Domain Understanding & Demo (10 pts): Explaining how the code works during Sunday demo walkthroughs.",
        "4. Integration & Autonomy (10 pts): Independent problem solving & proposal quality during Tarothons."
    ]
    for item in rubric_items:
        p_item = doc.add_paragraph(style='List Bullet')
        r_item = p_item.add_run(item)
        format_run(r_item, "Calibri", 10, color_rgb=(51, 51, 51))

    doc.add_paragraph()

    # Section 3: Squad Roles
    p_h4 = doc.add_paragraph()
    r_h4 = p_h4.add_run("🎭 SQUAD ROLES & THE CROSSOVER RULE")
    format_run(r_h4, "Segoe UI", 14, bold=True, color_rgb=(103, 58, 183))

    p_cross = doc.add_paragraph()
    r_cross = p_cross.add_run("🤝 THE CROSSOVER RULE: Your role is your primary domain lead, but cross-domain collaboration is 100% welcomed!")
    format_run(r_cross, "Calibri", 10.5, bold=True, italic=True, color_rgb=(103, 58, 183))

    p_hem = doc.add_paragraph()
    r_hem_t = p_hem.add_run("💸 HEMANTH — 'THE MONEY MAESTRO' (Finance Lead)\n")
    format_run(r_hem_t, "Segoe UI", 11, bold=True, color_rgb=(46, 125, 50))
    r_hem_b = p_hem.add_run("Playground: client/src/pages/PaymentSettings.jsx & api/payment/index.js")
    format_run(r_hem_b, "Calibri", 10)

    p_dhan = doc.add_paragraph()
    r_dhan_t = p_dhan.add_run("🎨 DHANUSHREE — 'THE VISUAL & VIBE QUEEN' (UI/UX Code Lead)\n")
    format_run(r_dhan_t, "Segoe UI", 11, bold=True, color_rgb=(156, 39, 176))
    r_dhan_b = p_dhan.add_run("Playground: client/src/pages/PublicCatalog.jsx, Landing.jsx, ExportCatalog.jsx & docs!")
    format_run(r_dhan_b, "Calibri", 10)

    p_dev = doc.add_paragraph()
    r_dev_t = p_dev.add_run("⚡ DEV ANAND — 'THE MASTERMIND ARCHITECT' (Tech Lead)\n")
    format_run(r_dev_t, "Segoe UI", 11, bold=True, color_rgb=(21, 101, 192))
    r_dev_b = p_dev.add_run("Playground: Voice STT engine, Gemini/Perplexity AI APIs (api/ai/), Shopify REST sync, and backend security.")
    format_run(r_dev_b, "Calibri", 10)

    doc.save(filepath)
    print(f"Updated: {filepath}")

# -------------------------------------------------------------
# DOCUMENT 2: TEAMMATE EVALUATION SHEET
# -------------------------------------------------------------
def build_evaluation_docx(filepath):
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("📝 THE TAROT CLUB — TEAMMATE EVALUATION SHEET")
    format_run(r_title, "Segoe UI", 18, bold=True, color_rgb=(103, 58, 183))

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("Evaluator: Dev Anand (Tech Lead)  |  Project: AI Catalog Agent\nRhythm: Logged every Sunday after the 15-minute weekly demo.")
    format_run(r_sub, "Calibri", 11, italic=True, color_rgb=(100, 100, 100))

    doc.add_paragraph()

    # Section 1: Rubric Table
    p_h1 = doc.add_paragraph()
    r_h1 = p_h1.add_run("🎯 40-POINT WEEKLY EVALUATION RUBRIC")
    format_run(r_h1, "Segoe UI", 14, bold=True, color_rgb=(103, 58, 183))

    rubric_data = [
        ["Criterion", "Max Points", "What to Check (Keep it Simple)"],
        ["1. Feature Completeness", "10 pts", "Is the weekly feature working and merged into the app?"],
        ["2. Code Quality & Mobile UX", "10 pts", "Clean React/Tailwind/Node code, responsive UI, no console errors."],
        ["3. Demo & Understanding", "10 pts", "Can they explain HOW their code works during the Sunday demo?"],
        ["4. Autonomy & Integration", "10 pts", "Solved blockers independently (or pitched a good Tarothon feature)."]
    ]

    tbl_rub = doc.add_table(rows=len(rubric_data), cols=3)
    tbl_rub.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_rub, color="D1C4E9", sz="6")
    rub_widths = [Inches(2.2), Inches(1.1), Inches(3.7)]

    for row_idx, row_content in enumerate(rubric_data):
        for col_idx, text in enumerate(row_content):
            cell = tbl_rub.cell(row_idx, col_idx)
            cell.width = rub_widths[col_idx]
            set_cell_margins(cell, top=80, bottom=80, left=80, right=80)
            p = cell.paragraphs[0]
            r = p.add_run(text)
            if row_idx == 0:
                set_cell_background(cell, "512DA8")
                format_run(r, "Segoe UI", 9.5, bold=True, color_rgb=(255, 255, 255))
            else:
                bg = "FAFAFA" if row_idx % 2 == 0 else "FFFFFF"
                set_cell_background(cell, bg)
                format_run(r, "Calibri", 9, color_rgb=(51, 51, 51))

    doc.add_paragraph()

    # Section 2: MASTER EVALUATION SCORE BOARD
    p_h2 = doc.add_paragraph()
    r_h2 = p_h2.add_run("📊 MASTER EVALUATION SCORE BOARD (EASY WEEKLY LOG)")
    format_run(r_h2, "Segoe UI", 14, bold=True, color_rgb=(103, 58, 183))

    master_scores = [
        ["Week", "Quest Stage", "💸 Hemanth Score", "🎨 Dhanushree Score", "Quick Lead Notes"],
        ["W1", "Core Storefront (UPI Settings & Grid)", "___ / 40", "___ / 40", ""],
        ["W2", "Payments & i18n (QR Gen & Keyboard)", "___ / 40", "___ / 40", ""],
        ["W3", "Exporters & Prices (Pricing DB & Cards)", "___ / 40", "___ / 40", ""],
        ["W4", "🃏 TAROTHON #1 (Innovation Feature)", "___ / 40", "___ / 40", ""],
        ["W5", "Dashboards (Toggles & Summary UI)", "___ / 40", "___ / 40", ""],
        ["W6", "Multimodal UI (Edit Fields & Camera)", "___ / 40", "___ / 40", ""],
        ["W7", "Landing & Polish (Security & Landing UI)", "___ / 40", "___ / 40", ""],
        ["W8", "🏆 TAROTHON #2 (Final Showcase)", "___ / 40", "___ / 40", ""],
        ["TOTAL", "8-WEEK CUMULATIVE SCORE", "___ / 320", "___ / 320", "Overall Progress"]
    ]

    tbl_master = doc.add_table(rows=len(master_scores), cols=5)
    tbl_master.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_master, color="D1C4E9", sz="6")
    m_widths = [Inches(0.6), Inches(2.5), Inches(1.2), Inches(1.3), Inches(1.4)]

    for r_i, r_c in enumerate(master_scores):
        for c_i, t_c in enumerate(r_c):
            cell = tbl_master.cell(r_i, c_i)
            cell.width = m_widths[c_i]
            set_cell_margins(cell, top=70, bottom=70, left=70, right=70)
            p_c = cell.paragraphs[0]
            if c_i in [2, 3]:
                p_c.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r_cell = p_c.add_run(t_c)
            if r_i == 0:
                set_cell_background(cell, "512DA8")
                format_run(r_cell, "Segoe UI", 9.5, bold=True, color_rgb=(255, 255, 255))
            elif r_i == len(master_scores) - 1:
                set_cell_background(cell, "D1C4E9")
                format_run(r_cell, "Segoe UI", 9.5, bold=True, color_rgb=(74, 20, 140))
            else:
                bg = "FAFAFA" if r_i % 2 == 0 else "FFFFFF"
                set_cell_background(cell, bg)
                format_run(r_cell, "Calibri", 9, color_rgb=(51, 51, 51))

    doc.add_paragraph()

    # Section 3: DETAILED WEEKLY BREAKDOWN LOGS
    p_h3 = doc.add_paragraph()
    r_h3 = p_h3.add_run("🗓️ DETAILED WEEKLY BREAKDOWN LOGS")
    format_run(r_h3, "Segoe UI", 14, bold=True, color_rgb=(103, 58, 183))

    weeks_info = [
        ("WEEK 1: Core Storefront Features", "UPI Settings & Bank Details (PaymentSettings.jsx)", "Public Catalog Layout (PublicCatalog.jsx)"),
        ("WEEK 2: Payments & Regional UI", "Dynamic QR Generator & WhatsApp Order link", "6-Language Regional Keyboard (LanguageSelector.jsx)"),
        ("WEEK 3: Exporters & Pricing Data", "Pricing Data Structures & API handlers", "Export Catalog UI Cards (ExportCatalog.jsx)"),
        ("WEEK 4: 🃏 TAROTHON #1 (INNOVATION HACKATHON)", "Finance Innovation Feature & Proposal", "UI/UX Innovation Feature & Proposal"),
        ("WEEK 5: Dashboards & Controls", "Payment Toggles & Summary UI", "Dashboard Layout & Search Filters"),
        ("WEEK 6: Multimodal UI & Product Edit", "Product Edit Pricing Fields", "Camera Capture Modal UI"),
        ("WEEK 7: Landing Page & Documentation Polish", "Payment Security Review", "Landing Page UI & Docs Polish"),
        ("WEEK 8: 🏆 TAROTHON #2 (FINAL SHOWCASE & LAUNCH)", "Final Showcase Innovation", "Final Showcase Innovation")
    ]

    for wk_title, hem_task, dhan_task in weeks_info:
        p_w = doc.add_paragraph()
        r_w = p_w.add_run(f"📌 {wk_title}")
        format_run(r_w, "Segoe UI", 11, bold=True, color_rgb=(74, 20, 140))

        log_data = [
            ["Teammate", "Assigned Task", "Complete (10)", "Quality (10)", "Demo (10)", "Autonomy (10)", "Total (40)"],
            ["💸 Hemanth", hem_task, "__ / 10", "__ / 10", "__ / 10", "__ / 10", "__ / 40"],
            ["🎨 Dhanushree", dhan_task, "__ / 10", "__ / 10", "__ / 10", "__ / 10", "__ / 40"]
        ]

        tbl_log = doc.add_table(rows=3, cols=7)
        tbl_log.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(tbl_log, color="E0E0E0", sz="4")
        log_widths = [Inches(1.2), Inches(2.3), Inches(0.7), Inches(0.7), Inches(0.7), Inches(0.7), Inches(0.7)]

        for r_i, r_c in enumerate(log_data):
            for c_i, t_c in enumerate(r_c):
                cell = tbl_log.cell(r_i, c_i)
                cell.width = log_widths[c_i]
                set_cell_margins(cell, top=50, bottom=50, left=50, right=50)
                p_c = cell.paragraphs[0]
                if c_i >= 2:
                    p_c.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r_cell = p_c.add_run(t_c)
                if r_i == 0:
                    set_cell_background(cell, "7E57C2")
                    format_run(r_cell, "Segoe UI", 8.5, bold=True, color_rgb=(255, 255, 255))
                else:
                    format_run(r_cell, "Calibri", 8.5, color_rgb=(51, 51, 51))

        p_fb = doc.add_paragraph()
        r_fb = p_fb.add_run("Lead Notes: _____________________________________________________________________\n")
        format_run(r_fb, "Calibri", 9, italic=True, color_rgb=(120, 120, 120))

    doc.save(filepath)
    print(f"Updated: {filepath}")

if __name__ == "__main__":
    base_dir = r"c:\Users\p3ace\OneDrive\Documents\New folder\AI-Catalog-Agent"
    doc1 = os.path.join(base_dir, "TEAM_ROADMAP_AND_LEARNING_PLAN.docx")
    doc2 = os.path.join(base_dir, "TEAMMATE_EVALUATION_SHEET.docx")
    
    build_roadmap_docx(doc1)
    build_evaluation_docx(doc2)
