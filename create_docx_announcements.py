import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

weeks_data = [
    {
        "week_num": 1,
        "title": "WEEK 1: CORE STOREFRONT",
        "dates": "2026-08-05 to 2026-08-12",
        "stage": "STOREFRONT FOUNDATION",
        "showcase": "NO",
        "filename": "Week_01_Briefing.docx",
        "tasks": [
            ("Hemanth", "Finance Lead", "UPI Settings & Bank Details Configuration", "PaymentSettings.jsx"),
            ("Dhanushree", "UI/UX Code Lead", "Public Catalog Layout & Product Grid System", "PublicCatalog.jsx"),
            ("Dev Anand", "Tech Lead", "Voice Input Engine & Express API Routes Setup", "routes/ai.js")
        ],
        "deliverables": [
            "UPI payment details form fully functional in seller settings.",
            "Public catalog grid rendering responsive product cards.",
            "Backend API route accepting voice STT input payloads.",
            "15-Minute Sunday Demo & Score Logging completed."
        ]
    },
    {
        "week_num": 2,
        "title": "WEEK 2: PAYMENTS & REGIONAL UI",
        "dates": "2026-08-13 to 2026-08-20",
        "stage": "PAYMENTS & I18N",
        "showcase": "NO",
        "filename": "Week_02_Briefing.docx",
        "tasks": [
            ("Hemanth", "Finance Lead", "Dynamic QR Generator & WhatsApp Order Link Integration", "PaymentHub.jsx"),
            ("Dhanushree", "UI/UX Code Lead", "6-Language Regional Keyboard & Language Selector UI", "LanguageContext.jsx"),
            ("Dev Anand", "Tech Lead", "AI Response Parsing & Supabase Database Handler", "services/db.js")
        ],
        "deliverables": [
            "Dynamic QR payment generator active for customer checkout.",
            "6-Language regional keyboard pills switching live UI text.",
            "AI JSON parsing connected directly to Supabase tables.",
            "15-Minute Sunday Demo & Score Logging completed."
        ]
    },
    {
        "week_num": 3,
        "title": "WEEK 3: EXPORTERS & PRICING",
        "dates": "2026-08-21 to 2026-08-28",
        "stage": "MARKETPLACE EXPORTS",
        "showcase": "NO",
        "filename": "Week_03_Briefing.docx",
        "tasks": [
            ("Hemanth", "Finance Lead", "Pricing Data Structures & Paise-Rupee Currency Converters", "routes/payment.js"),
            ("Dhanushree", "UI/UX Code Lead", "Export Catalog UI Cards (Shopify / Amazon / Meesho Export Modal)", "ExportModal.jsx"),
            ("Dev Anand", "Tech Lead", "Shopify REST Sync Endpoints & Variant Price Guardrails", "services/shopify.js")
        ],
        "deliverables": [
            "Paise-to-Rupee pricing converter active across checkout flows.",
            "Export modal supporting Shopify and marketplace format cards.",
            "Real-time Shopify REST sync handler pushing live inventory updates.",
            "15-Minute Sunday Demo & Score Logging completed."
        ]
    },
    {
        "week_num": 4,
        "title": "WEEK 4: MID SHOWCASE (TAROTHON #1)",
        "dates": "2026-08-29 to 2026-09-05",
        "stage": "INNOVATION HACKATHON",
        "showcase": "YES (MID REVIEW)",
        "filename": "Week_04_Briefing.docx",
        "tasks": [
            ("Hemanth", "Finance Lead", "Tarothon #1 Feature Pitch & Financial Flow Walkthrough", "TarothonPage.jsx"),
            ("Dhanushree", "UI/UX Code Lead", "Tarothon #1 UI Showcase & Live Component Demo Deck", "TarothonPage.jsx"),
            ("Dev Anand", "Tech Lead", "Technical Mentoring, Code Reviews & Core System Demo", "server/index.js")
        ],
        "deliverables": [
            "Live working prototype defense (Zero crash runtime).",
            "Tarothon #1 slide deck & feature walkthrough completed.",
            "Mid-Sprint Scorecard logged for all 3 members.",
            "Evaluator feedback compiled for Sprint Phase 2."
        ]
    },
    {
        "week_num": 5,
        "title": "WEEK 5: DASHBOARDS & CONTROLS",
        "dates": "2026-09-06 to 2026-09-13",
        "stage": "ANALYTICS & DASHBOARD",
        "showcase": "NO",
        "filename": "Week_05_Briefing.docx",
        "tasks": [
            ("Hemanth", "Finance Lead", "Payment Toggles & Inventory Revenue Summary UI", "PaymentSummary.jsx"),
            ("Dhanushree", "UI/UX Code Lead", "Dashboard Seller Panel & Catalog Status Filters", "Dashboard.jsx"),
            ("Dev Anand", "Tech Lead", "Voice Command Interpreter Microservice Endpoint", "routes/voice.js")
        ],
        "deliverables": [
            "Revenue summary cards rendering total inventory valuation in ₹.",
            "Seller dashboard panel supporting category filtering & quick actions.",
            "Voice command microservice accepting natural language edit instructions.",
            "15-Minute Sunday Demo & Score Logging completed."
        ]
    },
    {
        "week_num": 6,
        "title": "WEEK 6: MULTIMODAL UI & EDIT",
        "dates": "2026-09-14 to 2026-09-21",
        "stage": "MULTIMODAL EDITING",
        "showcase": "NO",
        "filename": "Week_06_Briefing.docx",
        "tasks": [
            ("Hemanth", "Finance Lead", "Product Edit Pricing Fields & Currency Sanitizer", "EditProduct.jsx"),
            ("Dhanushree", "UI/UX Code Lead", "Camera Capture Modal UI & Photo Upload Dropzone", "CameraModal.jsx"),
            ("Dev Anand", "Tech Lead", "Dual AI Fallback Handlers (Gemini 1.5 & Perplexity API)", "services/ai.js")
        ],
        "deliverables": [
            "Sanitized product edit modal supporting price & variant updates.",
            "Camera photo capture & direct image upload dropzone component.",
            "Dual-AI model failover pipeline preventing API timeouts.",
            "15-Minute Sunday Demo & Score Logging completed."
        ]
    },
    {
        "week_num": 7,
        "title": "WEEK 7: FINAL TAROTHON SHOWCASE",
        "dates": "2026-09-22 to 2026-09-29",
        "stage": "GRAND FINALE RELEASE",
        "showcase": "YES (GRAND FINALE)",
        "filename": "Week_07_Briefing.docx",
        "tasks": [
            ("Hemanth", "Finance Lead", "Grand Finale Presentation & Financial Flow Walkthrough", "FinalDemo.jsx"),
            ("Dhanushree", "UI/UX Code Lead", "Grand Finale UI Showcase & Responsive Layout Defense", "FinalDemo.jsx"),
            ("Dev Anand", "Tech Lead", "Final Production Release & Architecture Defense", "PROJECT_SHOWCASE.md")
        ],
        "deliverables": [
            "Live production release deployed on Vercel with clean build.",
            "Grand Tarothon presentation pitch & architecture defense.",
            "100% Cumulative 7-Week Evaluation Scores logged.",
            "Final production documentation signed off."
        ]
    }
]

def set_cell_background(cell, hex_color):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

for w in weeks_data:
    doc = Document()

    # Document margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Header Title Banner
    header_table = doc.add_table(rows=1, cols=1)
    header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = header_table.cell(0, 0)
    set_cell_background(cell, "0F172A")
    set_cell_margins(cell, top=200, bottom=200, left=200, right=200)

    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"📢 {w['title']}")
    run.font.name = 'Calibri'
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = RGBColor(56, 189, 248) # Cyan

    p2 = cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run2 = p2.add_run("AI CATALOG AGENT — TEAM SPRINT ANNOUNCEMENT")
    run2.font.name = 'Calibri'
    run2.font.size = Pt(12)
    run2.font.bold = True
    run2.font.color.rgb = RGBColor(148, 163, 184) # Muted gray

    p3 = cell.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run3 = p3.add_run(f"📅 {w['dates']}   |   ⚙️ STAGE: {w['stage']}   |   🌟 SHOWCASE: {w['showcase']}")
    run3.font.name = 'Calibri'
    run3.font.size = Pt(10)
    run3.font.bold = True
    run3.font.color.rgb = RGBColor(253, 224, 71) # Gold

    doc.add_paragraph()

    # Section 1: Team Member Roles & Responsibilities
    h = doc.add_heading("👥 TEAM MEMBER ROLES & SPECIFIC RESPONSIBILITIES", level=2)
    h.runs[0].font.name = 'Calibri'
    h.runs[0].font.color.rgb = RGBColor(30, 64, 175)

    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table.rows[0].cells
    headers = ["Team Member", "Role", "Assigned Task", "Target File"]
    widths = [Inches(1.3), Inches(1.5), Inches(2.7), Inches(1.3)]

    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], "1E3A8A")
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=150, right=150)
        p = hdr_cells[i].paragraphs[0]
        p.runs[0].font.name = 'Calibri'
        p.runs[0].font.bold = True
        p.runs[0].font.size = Pt(10)
        p.runs[0].font.color.rgb = RGBColor(255, 255, 255)

    for idx, (member, role, task, target_file) in enumerate(w['tasks']):
        row_cells = table.add_row().cells
        bg = "F8FAFC" if idx % 2 == 1 else "FFFFFF"
        for i, val in enumerate([member, role, task, target_file]):
            row_cells[i].text = val
            set_cell_background(row_cells[i], bg)
            set_cell_margins(row_cells[i], top=100, bottom=100, left=120, right=120)
            p = row_cells[i].paragraphs[0]
            p.runs[0].font.name = 'Calibri'
            p.runs[0].font.size = Pt(10.5)
            if i == 0:
                p.runs[0].font.bold = True

    doc.add_paragraph()

    # Section 2: Deliverables
    h2 = doc.add_heading("📋 WEEKLY DELIVERABLES", level=2)
    h2.runs[0].font.name = 'Calibri'
    h2.runs[0].font.color.rgb = RGBColor(30, 64, 175)

    for item in w['deliverables']:
        p = doc.add_paragraph(style='List Bullet')
        run = p.add_run(item)
        run.font.name = 'Calibri'
        run.font.size = Pt(11)

    doc.add_paragraph()

    # Section 3: Evaluation Rubric
    h3 = doc.add_heading("⭐ 40-POINT EVALUATION RUBRIC", level=2)
    h3.runs[0].font.name = 'Calibri'
    h3.runs[0].font.color.rgb = RGBColor(30, 64, 175)

    rubric_table = doc.add_table(rows=1, cols=1)
    rubric_cell = rubric_table.cell(0, 0)
    set_cell_background(rubric_cell, "F0FDF4")
    set_cell_margins(rubric_cell, top=150, bottom=150, left=150, right=150)

    rubric_items = [
        "Task Ownership (10 pts): Read, tested, and contributed to assigned features.",
        "Code Quality & Mobile UX (10 pts): Clean code, responsive UI, zero console errors.",
        "Demo & Understanding (10 pts): Clear explanation during Sunday demo.",
        "Autonomy & Integration (10 pts): Solved blockers independently without delays."
    ]

    for idx, item in enumerate(rubric_items):
        p = rubric_cell.paragraphs[0] if idx == 0 else rubric_cell.add_paragraph()
        run = p.add_run(f"•  {item}")
        run.font.name = 'Calibri'
        run.font.size = Pt(10.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(22, 101, 52)

    output_path = os.path.join(r"c:\Users\p3ace\OneDrive\Documents\New folder\AI-Catalog-Agent", w['filename'])
    doc.save(output_path)
    print(f"Generated: {output_path}")

print("All DOCX files created successfully!")
